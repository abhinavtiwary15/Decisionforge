#!/usr/bin/env python3
"""
pipeline/ingest_purchase_register.py
====================================
Flexible Purchase Register Ingestion Module.

Accepts any CSV or Excel upload, detects headers, and processes data using
client-defined column mappings. Normalises columns, dates, numerics, and validates GSTINs.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import sys
from datetime import datetime
from typing import Any, Optional

# ── UTF-8 console output (must be before any print) ─────────────────────────
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

# ── Validator import — supports both script and module usage ─────────────────
_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJ_DIR = os.path.dirname(_THIS_DIR)
if _PROJ_DIR not in sys.path:
    sys.path.insert(0, _PROJ_DIR)
try:
    from pipeline.validators import validate_gstin
except ImportError:
    from validators import validate_gstin  # type: ignore[no-redef]


# ═══════════════════════════════════════════════════════════════════════════════
# File Reading Helpers
# ═══════════════════════════════════════════════════════════════════════════════

def detect_file_headers(filepath: str) -> list[str]:
    """Reads the first row of a CSV or Excel file and returns its headers."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".json":
        # Handle simple JSON list of dicts for convenience/testing if needed, but not requested.
        pass
    if ext in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError as exc:
            raise ImportError("openpyxl is required to read Excel files.") from exc
        
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        if not wb.sheetnames:
            wb.close()
            raise ValueError("Excel file has no sheets.")
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration:
            wb.close()
            raise ValueError("Excel file is empty.")
        wb.close()
        return [str(h).strip() for h in headers if h is not None]
    else:
        # Default to CSV
        for encoding in ("utf-8", "utf-8-sig", "latin1", "cp1252"):
            try:
                with open(filepath, "r", encoding=encoding) as fh:
                    reader = csv.reader(fh)
                    headers = next(reader)
                    return [h.strip() for h in headers if h is not None]
            except (UnicodeDecodeError, StopIteration):
                continue
        raise ValueError("Could not read file headers as CSV.")


def read_file_rows(filepath: str) -> list[dict[str, Any]]:
    """Reads CSV or Excel file and returns a list of dictionaries mapping header to value."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError as exc:
            raise ImportError("openpyxl is required to read Excel files.") from exc
        
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [str(h).strip() for h in next(rows_iter) if h is not None]
        except StopIteration:
            wb.close()
            return []
        
        rows = []
        for row_values in rows_iter:
            if all(v is None or (isinstance(v, str) and v.strip() == "") for v in row_values):
                continue
            # Match headers and values
            row_dict = {}
            for idx, val in enumerate(row_values):
                if idx < len(headers):
                    row_dict[headers[idx]] = val
            rows.append(row_dict)
        wb.close()
        return rows
    else:
        # Default to CSV
        for encoding in ("utf-8", "utf-8-sig", "latin1", "cp1252"):
            try:
                with open(filepath, "r", encoding=encoding) as fh:
                    reader = csv.DictReader(fh)
                    rows = []
                    for row in reader:
                        # Clean whitespace from keys and values
                        cleaned = {k.strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k is not None}
                        # Skip empty rows
                        if all(v is None or v == "" for v in cleaned.values()):
                            continue
                        rows.append(cleaned)
                    return rows
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not parse file rows as CSV.")


# ═══════════════════════════════════════════════════════════════════════════════
# Processing & Validation Logic
# ═══════════════════════════════════════════════════════════════════════════════

def parse_date(val: Any) -> tuple[Optional[str], Optional[str]]:
    """
    Parses date trying DD-MM-YYYY, DD/MM/YYYY, YYYY-MM-DD formats.
    Returns (normalized_date_str_or_none, error_msg_or_none).
    """
    if val is None or (isinstance(val, str) and val.strip() == ""):
        return None, "Date is missing"
    raw = str(val).strip()
    
    # Try parsing different formats
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            parsed = datetime.strptime(raw, fmt)
            return parsed.strftime("%Y-%m-%d"), None
        except ValueError:
            continue
            
    return None, f"Date '{raw}' is in an unsupported format. Must be DD-MM-YYYY, DD/MM/YYYY, or YYYY-MM-DD."


def parse_float(val: Any, field_name: str) -> tuple[float, Optional[str]]:
    """Parses value as float. Returns (float_val, error_msg_or_none)."""
    if val is None or (isinstance(val, str) and val.strip() == ""):
        return 0.0, None
    raw = str(val).replace(",", "").strip()
    try:
        return float(raw), None
    except ValueError:
        return 0.0, f"Field '{field_name}' has non-numeric value '{val}'."


def ingest_purchase_register(
    filepath: str,
    mapping: dict[str, str]
) -> dict[str, Any]:
    """
    Ingests and normalises the purchase register using the provided mapping.
    
    Parameters:
      filepath: path to file (CSV or Excel)
      mapping: dictionary of {our_field: their_column_name}
      
    Returns:
      {
         "records": list of normalised valid rows,
         "data_quality_flags": list of invalid GSTIN rows,
         "summary": {
             "total": int,
             "clean": int,
             "flagged": int,
             "issues_count": int
         },
         "issues": list of dicts {row: int, invoice_number: str, field: str, problem: str}
      }
    """
    rows = read_file_rows(filepath)
    records = []
    data_quality_flags = []
    issues = []
    
    # Required fields in mapping
    required_mapped = ["invoice_number", "vendor_gstin", "taxable_value"]
    
    for idx, row in enumerate(rows, start=2): # Start at 2 (assuming header is row 1)
        invoice_number = None
        # Extract invoice_number first for reporting reference
        inv_col = mapping.get("invoice_number")
        if inv_col and inv_col in row:
            invoice_number = str(row[inv_col] or "").strip()
        
        # Check mapping required presence in row
        row_missing = []
        for req in required_mapped:
            mapped_col = mapping.get(req)
            if not mapped_col or mapped_col not in row or row[mapped_col] is None or str(row[mapped_col]).strip() == "":
                row_missing.append(req)
                
        if row_missing:
            issues.append({
                "row": idx,
                "invoice_number": invoice_number or "<unknown>",
                "field": "mapping",
                "problem": f"Row is missing required columns: {', '.join(row_missing)}"
            })
            continue

        # Extract values
        raw_inv_num = str(row[mapping["invoice_number"]] or "").strip()
        raw_gstin = str(row[mapping["vendor_gstin"]] or "").strip()
        raw_taxable = row[mapping["taxable_value"]]
        
        raw_name = row.get(mapping.get("vendor_name")) if mapping.get("vendor_name") else None
        raw_date = row.get(mapping.get("invoice_date")) if mapping.get("invoice_date") else None
        raw_cgst = row.get(mapping.get("cgst")) if mapping.get("cgst") else None
        raw_sgst = row.get(mapping.get("sgst")) if mapping.get("sgst") else None
        raw_igst = row.get(mapping.get("igst")) if mapping.get("igst") else None
        raw_itc = row.get(mapping.get("total_itc_claimed")) if mapping.get("total_itc_claimed") else None

        # Parse numeric taxable value
        taxable_val, err_tax = parse_float(raw_taxable, "taxable_value")
        if err_tax:
            issues.append({
                "row": idx,
                "invoice_number": raw_inv_num,
                "field": "taxable_value",
                "problem": err_tax
            })
            continue

        # Parse optional fields
        err_list = []
        
        # Parse Date if present
        parsed_date = None
        if raw_date not in (None, ""):
            parsed_date, err_date = parse_date(raw_date)
            if err_date:
                err_list.append(("invoice_date", err_date))
        
        # Parse taxes
        cgst_val, err_cgst = parse_float(raw_cgst, "cgst")
        if err_cgst: err_list.append(("cgst", err_cgst))
        
        sgst_val, err_sgst = parse_float(raw_sgst, "sgst")
        if err_sgst: err_list.append(("sgst", err_sgst))
        
        igst_val, err_igst = parse_float(raw_igst, "igst")
        if err_igst: err_list.append(("igst", err_igst))
        
        itc_val, err_itc = parse_float(raw_itc, "total_itc_claimed")
        if err_itc: err_list.append(("total_itc_claimed", err_itc))
        
        if not mapping.get("total_itc_claimed") and not err_cgst and not err_sgst and not err_igst:
            # Derive total_itc if not mapped
            itc_val = round(cgst_val + sgst_val + igst_val, 2)

        # If there are parsing issues, log them and skip clean record
        if err_list:
            for fld, err_msg in err_list:
                issues.append({
                    "row": idx,
                    "invoice_number": raw_inv_num,
                    "field": fld,
                    "problem": err_msg
                })
            continue

        # Validate GSTIN
        valid_gstin, err_gstin = validate_gstin(raw_gstin)
        
        record_dict = {
            "invoice_number": raw_inv_num,
            "vendor_gstin": raw_gstin.strip().upper(),
            "vendor_name": str(raw_name).strip() if raw_name is not None else None,
            "invoice_date": parsed_date,
            "taxable_value": taxable_val,
            "cgst": cgst_val,
            "sgst": sgst_val,
            "igst": igst_val,
            "total_itc_claimed": itc_val
        }
        
        if not valid_gstin:
            # Log as data quality flag
            data_quality_flags.append({
                "vendor_gstin": raw_gstin,
                "vendor_name": record_dict["vendor_name"],
                "invoice_number": raw_inv_num,
                "invoice_date": parsed_date,
                "flag_type": "INVALID_GSTIN",
                "flag_detail": err_gstin
            })
            issues.append({
                "row": idx,
                "invoice_number": raw_inv_num,
                "field": "vendor_gstin",
                "problem": err_gstin
            })
        else:
            records.append(record_dict)

    skipped = len(rows) - len(records) - len(data_quality_flags)
    summary = {
        "total": len(rows),
        "clean": len(records),
        "flagged": len(data_quality_flags),
        "skipped": skipped,
        "issues_count": len(issues)
    }

    return {
        "records": records,
        "data_quality_flags": data_quality_flags,
        "summary": summary,
        "issues": issues
    }


# ═══════════════════════════════════════════════════════════════════════════════
# Test Suite
# ═══════════════════════════════════════════════════════════════════════════════

def run_test_suite():
    """Runs the internal ingestion validation tests."""
    print("=" * 70)
    print("RUNNING PURCHASE REGISTER INGESTION TESTS")
    print("=" * 70)
    
    # 1. Create S1 CSV (Tally format) and S2 Excel (Zoho format) in memory/temp files
    import tempfile
    
    # Standard valid GSTINs
    GSTIN_1 = "27AABCU9603R1ZV"
    GSTIN_2 = "29AADCB2230M1ZV"
    GSTIN_INVALID = "99AABCU9603R1ZV"
    
    s1_content = (
        "Date,Invoice No,Party Name,Party GSTIN,Taxable Amt,CGST Amt,SGST Amt,IGST Amt,Total ITC\n"
        f"05/04/2026,INV-001,Reliance Pvt, {GSTIN_1}, 10000.0, 900.0, 900.0, 0.0, 1800.0\n"
        f"12-04-2026,INV-002,Infosys Ltd,  {GSTIN_2}, 20000.0, 0.0, 0.0, 3600.0, 3600.0\n"
        f"15/04/2026,INV-003,Bad Supplier, {GSTIN_INVALID}, 5000.0, 450.0, 450.0, 0.0, 900.0\n" # Invalid GSTIN
        f"32-13-2026,INV-004,Date Error,    {GSTIN_1}, 15000.0, 1350.0, 1350.0, 0.0, 2700.0\n" # Bad Date
    )
    
    # Let's write S1 to a temp file
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f_s1:
        f_s1.write(s1_content)
        s1_path = f_s1.name
        
    # Let's create S2 Zoho Excel in memory and save
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        # Zoho columns in different order & naming
        headers = ["Invoice Number", "Invoice Date", "GSTIN of Vendor", "Vendor Name", "Taxable Value", "CGST", "SGST", "IGST", "Total Claimed"]
        ws.append(headers)
        ws.append(["INV-001", "2026-04-05", GSTIN_1, "Reliance Pvt", 10000.0, 900.0, 900.0, 0.0, 1800.0])
        ws.append(["INV-002", "12/04/2026", GSTIN_2, "Infosys Ltd", 20000.0, 0.0, 0.0, 3600.0, 3600.0])
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f_s2:
            wb.save(f_s2.name)
            s2_path = f_s2.name
            wb.close()
            excel_ok = True
    except Exception as e:
        print(f"Skipping Excel tests because of setup issues: {e}")
        excel_ok = False
        s2_path = None

    # Define mappings
    s1_mapping = {
        "invoice_number": "Invoice No",
        "vendor_gstin": "Party GSTIN",
        "vendor_name": "Party Name",
        "invoice_date": "Date",
        "taxable_value": "Taxable Amt",
        "cgst": "CGST Amt",
        "sgst": "SGST Amt",
        "igst": "IGST Amt",
        "total_itc_claimed": "Total ITC"
    }
    
    s2_mapping = {
        "invoice_number": "Invoice Number",
        "vendor_gstin": "GSTIN of Vendor",
        "vendor_name": "Vendor Name",
        "invoice_date": "Invoice Date",
        "taxable_value": "Taxable Value",
        "cgst": "CGST",
        "sgst": "SGST",
        "igst": "IGST",
        "total_itc_claimed": "Total Claimed"
    }

    results = []
    
    try:
        # Test S1 Ingestion
        print("\n--- Test S1: CSV Tally-style columns & parsing ---")
        res1 = ingest_purchase_register(s1_path, s1_mapping)
        # Should parse 4 rows:
        # Row 1 (INV-001) - Clean
        # Row 2 (INV-002) - Clean
        # Row 3 (INV-003) - Invalid GSTIN flagged (under data_quality_flags, and also in issues)
        # Row 4 (INV-004) - Bad Date issue (in issues)
        
        clean_invs = [r["invoice_number"] for r in res1["records"]]
        flagged_invs = [r["invoice_number"] for r in res1["data_quality_flags"]]
        issue_invs = [i["invoice_number"] for i in res1["issues"]]
        
        print(f"Clean: {clean_invs}")
        print(f"Flagged (Invalid GSTIN): {flagged_invs}")
        print(f"Issues (Date/GSTIN errors): {res1['issues']}")
        
        assert "INV-001" in clean_invs
        assert "INV-002" in clean_invs
        assert "INV-003" in flagged_invs
        assert any(i["invoice_number"] == "INV-004" and i["field"] == "invoice_date" for i in res1["issues"])
        assert any(i["invoice_number"] == "INV-003" and i["field"] == "vendor_gstin" for i in res1["issues"])
        
        results.append(("S1: Ingestion & validations", "PASSED"))
    except Exception as e:
        results.append(("S1: Ingestion & validations", f"FAILED: {e}"))

    if excel_ok and s2_path:
        try:
            # Test S2 Ingestion
            print("\n--- Test S2: Excel Zoho-style columns ---")
            res2 = ingest_purchase_register(s2_path, s2_mapping)
            clean_invs2 = [r["invoice_number"] for r in res2["records"]]
            print(f"Clean: {clean_invs2}")
            
            assert "INV-001" in clean_invs2
            assert "INV-002" in clean_invs2
            assert len(res2["issues"]) == 0
            
            results.append(("S2: Excel Zoho mapping", "PASSED"))
        except Exception as e:
            results.append(("S2: Excel Zoho mapping", f"FAILED: {e}"))
            
    # Cleanup temp files
    try:
        os.unlink(s1_path)
        if s2_path:
            os.unlink(s2_path)
    except:
        pass
        
    print("\n" + "=" * 70)
    print("TEST SUITE RESULTS")
    print("=" * 70)
    for name, status in results:
        print(f"  [{status}] {name}")
    print("=" * 70)
    
    if any(status.startswith("FAILED") for name, status in results):
        sys.exit(1)


# ═══════════════════════════════════════════════════════════════════════════════
# Main Entry Point
# ═══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    import argparse
    
    parser = argparse.ArgumentParser(description="Flexible Ingest module for Purchase Register.")
    parser.add_argument("--detect-columns", help="Detect column headers from file")
    parser.add_argument("--ingest", help="Ingest file using column mapping")
    parser.add_argument("--mapping", help="JSON string representing the column mapping mapping")
    parser.add_argument("--run-tests", action="store_true", help="Run the test suite")
    
    args = parser.parse_args()
    
    if args.run_tests:
        run_test_suite()
    elif args.detect_columns:
        try:
            hdrs = detect_file_headers(args.detect_columns)
            print(json.dumps(hdrs))
        except Exception as e:
            print(json.dumps({"error": str(e)}))
            sys.exit(1)
    elif args.ingest:
        if not args.mapping:
            print(json.dumps({"error": "The --mapping parameter is required for ingestion."}))
            sys.exit(1)
        try:
            mapping_dict = json.loads(args.mapping)
            res = ingest_purchase_register(args.ingest, mapping_dict)
            print(json.dumps(res))
        except Exception as e:
            print(json.dumps({"error": str(e)}))
            sys.exit(1)
    else:
        parser.print_help()
